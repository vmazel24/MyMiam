"""Import official Ciqual 2025 XLSX, preserving unknowns and detection limits."""
import argparse
import json
import sys
from pathlib import Path
sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from openpyxl import load_workbook
from mymiam.storage import Store
from mymiam.nutrition import normalize

URL = "https://ciqual.anses.fr/cms/sites/default/files/inline-files/Table%20Ciqual%202025_FR_2025_11_03.xlsx"


def cell(value):
    text = str(value if value is not None else "").strip().replace(",", ".")
    if text in ("", "-", "None"):
        return None, "Valeur absente"
    if text.lower() == "traces":
        return None, "Traces : quantité non chiffrée"
    if text.startswith("<"):
        # Preserve uncertainty rather than replacing a detection limit with zero.
        return None, "Inférieur à " + text[1:].strip()
    try:
        return float(text), None
    except ValueError:
        return None, "Valeur non chiffrée"


def run(path, directory):
    store = Store(directory)
    book = load_workbook(path, read_only=True, data_only=True)
    rows = book["composition nutritionnelle"].iter_rows(values_only=True)
    headers = [" ".join(str(v).split()) for v in next(rows)]
    columns = {"kcal": 10, "protein": 14, "carbs": 16, "fat": 17, "fiber": 26}
    assert "kcal" in headers[10] and "Jones" in headers[14] and "Glucides" in headers[16]
    count = 0
    with store.connect() as db:
        for row in rows:
            if not row[6] or not row[7]:
                continue
            nutrients, flags = {}, {}
            for key, index in columns.items():
                nutrients[key], flag = cell(row[index])
                if flag:
                    flags[key] = flag
            db.execute("INSERT OR REPLACE INTO foods VALUES (?,?,?,?,?,?)", (
                "ciqual:" + str(row[6]), row[7], normalize(row[7]), "Ciqual 2025 · Anses",
                json.dumps(nutrients), json.dumps(flags)))
            count += 1
    book.close()
    store.set_setting("ciqual", {"version": "2025", "foods": count, "url": URL,
                                "license": "Licence Ouverte 2.0"})
    print(f"{count} aliments Ciqual 2025 importés")


if __name__ == "__main__":
    parser = argparse.ArgumentParser()
    parser.add_argument("xlsx")
    parser.add_argument("--instance", default="instance")
    args = parser.parse_args()
    run(args.xlsx, args.instance)
